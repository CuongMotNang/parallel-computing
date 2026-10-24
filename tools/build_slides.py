# -*- coding: utf-8 -*-
"""Sinh slide: Con trỏ trong C/C++ - Báo cáo giữa kỳ môn Tính toán song song.

Chạy:    python tools/build_slides.py
Kết quả: slides/ConTro_C_CPP.pptx
"""
import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

# ---------------- Màu & font ----------------
BLUE = RGBColor(0x1F, 0x4E, 0x79)        # xanh đậm: thanh tiêu đề
BLUE_LIGHT = RGBColor(0x2E, 0x75, 0xB6)  # xanh nhấn
BLUE_PALE = RGBColor(0xDE, 0xEB, 0xF7)   # nền ô nhạt
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x59, 0x59, 0x59)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CODE_BG = RGBColor(0xF4, 0xF6, 0xF8)
CODE_BORDER = RGBColor(0xBF, 0xBF, 0xBF)
COMMENT = RGBColor(0x00, 0x80, 0x00)
ORANGE = RGBColor(0xC5, 0x5A, 0x11)
ORANGE_PALE = RGBColor(0xFC, 0xE4, 0xD6)
RED = RGBColor(0xC0, 0x00, 0x00)
GREEN = RGBColor(0x00, 0x80, 0x00)
FONT, MONO = "Calibri", "Consolas"

TOTAL = 13
FOOTER = "Tính toán song song – Báo cáo giữa kỳ – Con trỏ trong C/C++"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
I = Inches


# ---------------- Tiện ích ----------------
def run(p, text, size=20, bold=False, color=BLACK, font=FONT):
    r = p.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold = font, Pt(size), bold
    f.color.rgb = color
    return r


def rich(p, text, size=20, color=BLACK, bold=False):
    """`code` -> Consolas xanh, **đậm** -> in đậm."""
    for tok in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not tok:
            continue
        if tok.startswith("`"):
            run(p, tok[1:-1], size, bold, BLUE, MONO)
        elif tok.startswith("**"):
            run(p, tok[2:-2], size, True, color)
        else:
            run(p, tok, size, bold, color)


def rect(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    sh = s.shapes.add_shape(shape, x, y, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1)
    sh.shadow.inherit = False
    return sh


def frame_lines(tf, lines, size, color=BLACK, bold=False, align=PP_ALIGN.LEFT,
                font=FONT, space=0):
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if space:
            p.space_after = Pt(space)
        if font == MONO:
            run(p, line, size, bold, color, MONO)
        else:
            rich(p, line, size, color, bold)


def text(s, x, y, w, h, t, size=20, bold=False, color=BLACK,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    frame_lines(tf, t.split("\n"), size, color, bold, align, font)
    return tb


def label(sh, t, size=16, bold=False, color=BLACK, font=FONT, align=PP_ALIGN.CENTER):
    """Ghi chữ vào trong một hình."""
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = I(0.05)
    tf.margin_top = tf.margin_bottom = I(0.03)
    frame_lines(tf, t.split("\n"), size, color, bold, align, font)


def bullets(s, x, y, w, h, items, size=22, space=10):
    """items: chuỗi (cấp 1) hoặc ('-', chuỗi) (cấp 2)."""
    tf = s.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        if isinstance(it, tuple):
            run(p, "      –  ", size - 2, True, BLUE_LIGHT)
            rich(p, it[1], size - 2)
        else:
            run(p, "■  ", size - 8, False, BLUE_LIGHT)
            rich(p, it, size)


def code(s, x, y, w, h, src, size=16):
    box = rect(s, x, y, w, h, CODE_BG, line=CODE_BORDER)
    rect(s, x, y, I(0.08), h, BLUE_LIGHT)          # vạch xanh bên trái
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left, tf.margin_right = I(0.25), I(0.1)
    for i, line in enumerate(src.strip("\n").split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        k = line.find("//")
        body, cmt = (line, "") if k < 0 else (line[:k], line[k:])
        if body or not cmt:
            run(p, body or " ", size, font=MONO)
        if cmt:
            run(p, cmt, size, color=COMMENT, font=MONO)


def table(s, x, y, w, rows, col_w, size=16, row_h=0.45):
    tbl = s.shapes.add_table(len(rows), len(rows[0]), x, y, w,
                             I(row_h * len(rows))).table
    for i, cw in enumerate(col_w):
        tbl.columns[i].width = I(cw)
    for r, row in enumerate(rows):
        tbl.rows[r].height = I(row_h)
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if r == 0 else (BLUE_PALE if r % 2 == 0 else WHITE)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            if r == 0:
                p.alignment = PP_ALIGN.CENTER
                run(p, val, size, True, WHITE)
            else:
                rich(p, val, size)


def arrow(s, x1, y1, x2, y2, color=BLUE_LIGHT):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(2.25)
    ln = c.line._get_or_add_ln()
    ln.append(ln.makeelement(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
        {"type": "triangle", "w": "med", "len": "med"}))


def slide(title, num):
    """Slide nội dung: thanh tiêu đề xanh, chân trang, số trang."""
    s = prs.slides.add_slide(BLANK)
    bar = rect(s, 0, 0, SW, I(1.05), BLUE)
    label(bar, title, 32, True, WHITE, align=PP_ALIGN.LEFT)
    bar.text_frame.margin_left = I(0.5)
    rect(s, 0, I(1.05), SW, I(0.06), BLUE_LIGHT)
    rect(s, I(0.5), I(7.0), SW - I(1.0), Pt(1), BLUE_LIGHT)
    text(s, I(0.5), I(7.02), I(9), I(0.4), FOOTER, 12, color=GRAY)
    text(s, SW - I(2.5), I(7.02), I(2.0), I(0.4), f"{num} / {TOTAL}", 12,
         color=GRAY, align=PP_ALIGN.RIGHT)
    return s


def note(s, t):
    s.notes_slide.notes_text_frame.text = t


L, TOP, FULL_W = I(0.6), I(1.45), SW - I(1.2)

# ===== Slide 1 – Trang bìa =====
s = prs.slides.add_slide(BLANK)
band = rect(s, 0, 0, SW, I(1.2), BLUE)
label(band, "TRƯỜNG ……………………  •  KHOA ……………………", 18, color=WHITE)
rect(s, 0, I(1.2), SW, I(0.06), BLUE_LIGHT)
text(s, 0, I(1.65), SW, I(0.5), "MÔN HỌC: TÍNH TOÁN SONG SONG", 22, True,
     BLUE_LIGHT, PP_ALIGN.CENTER)
text(s, 0, I(2.15), SW, I(0.5), "BÁO CÁO GIỮA KỲ", 20, color=GRAY, align=PP_ALIGN.CENTER)
text(s, 0, I(2.7), SW, I(1.0), "CON TRỎ (POINTER) TRONG C/C++", 44, True, BLUE,
     PP_ALIGN.CENTER)
rect(s, I(5.17), I(3.75), I(3.0), Pt(3), BLUE_LIGHT)
text(s, I(3.67), I(4.0), I(6), I(0.4), "Nhóm thực hiện", 18, True, align=PP_ALIGN.CENTER)
table(s, I(3.67), I(4.45), I(6.0), [
    ["STT", "Họ và tên", "Mã sinh viên"],
    ["1", "Bùi Xuân Chung", ""],
    ["2", "Lê Văn Cường", ""],
    ["3", "Đào Ngọc Cường", ""],
], [0.8, 3.0, 2.2], 16, 0.42)
text(s, 0, I(6.4), SW, I(0.4), "Giảng viên hướng dẫn: ……………………", 16,
     color=GRAY, align=PP_ALIGN.CENTER)
rect(s, 0, SH - I(0.15), SW, I(0.15), BLUE)
note(s, "Giới thiệu nhóm, môn học, chủ đề. Nhớ điền mã sinh viên, tên trường/khoa, giảng viên.")

# ===== Slide 2 – Nội dung =====
s = slide("Nội dung trình bày", 2)
agenda = ["Bộ nhớ và địa chỉ", "Khái niệm con trỏ", "Con trỏ và mảng",
          "Con trỏ và hàm", "Cấp phát bộ nhớ động", "Lỗi thường gặp",
          "Con trỏ trong C++ hiện đại", "Con trỏ trong tính toán song song",
          "Ví dụ minh họa & Kết luận"]
for i, item in enumerate(agenda):
    col, row = divmod(i, 5)
    x, y = I(1.2 + col * 5.8), I(1.65 + row * 1.0)
    c = rect(s, x, y, I(0.65), I(0.65), BLUE if i < 7 else BLUE_LIGHT, shape=MSO_SHAPE.OVAL)
    label(c, str(i + 1), 20, True, WHITE)
    text(s, x + I(0.9), y + I(0.03), I(4.8), I(0.6), item, 24,
         bold=(i == 7), anchor=MSO_ANCHOR.MIDDLE)
note(s, "Nhấn mạnh phần 8 gắn với môn Tính toán song song.")

# ===== Slide 3 – Bộ nhớ và địa chỉ =====
s = slide("1. Bộ nhớ và địa chỉ", 3)
bullets(s, L, TOP, I(6.6), I(3.0), [
    "RAM gồm các ô nhớ 1 byte, mỗi ô có một **địa chỉ** riêng",
    "Khai báo biến → được cấp vùng nhớ (`int` thường 4 byte)",
    "Mỗi biến có: **tên, kiểu, giá trị, địa chỉ**",
    "Toán tử `&` lấy địa chỉ của biến",
])
code(s, L, I(4.55), I(6.6), I(1.6), """
int x = 10;
printf("Gia tri: %d\\n", x);
printf("Dia chi: %p\\n", (void*)&x);
""", 17)
text(s, I(7.8), TOP, I(5), I(0.4), "Minh họa bộ nhớ", 18, True, BLUE, PP_ALIGN.CENTER)
for i, (ad, v) in enumerate(zip(["0x1000", "0x1004", "0x1008", "0x100C"],
                                ["…", "10", "…", "…"])):
    y = I(2.05 + i * 0.95)
    b = rect(s, I(9.3), y, I(2.2), I(0.85), BLUE_PALE if i == 1 else WHITE, line=BLUE)
    label(b, v, 22, i == 1)
    text(s, I(7.9), y + I(0.2), I(1.3), I(0.5), ad, 16, color=GRAY, font=MONO,
         align=PP_ALIGN.RIGHT)
text(s, I(11.6), I(3.15), I(1.3), I(0.5), "← x", 22, True, BLUE)
text(s, I(7.8), I(5.95), I(5), I(0.5), "&x  →  0x1004", 18, color=GRAY, font=MONO,
     align=PP_ALIGN.CENTER)
note(s, "Địa chỉ in ra mỗi lần chạy có thể khác nhau do ASLR.")

# ===== Slide 4 – Con trỏ là gì =====
s = slide("2. Con trỏ là gì?", 4)
bullets(s, L, TOP, I(6.6), I(2.8), [
    "**Con trỏ** là biến lưu **địa chỉ** của biến khác",
    "Khai báo: `kiểu *tên;`",
    "`&x` : lấy địa chỉ của x",
    "`*p` : truy cập giá trị tại địa chỉ p (dereference)",
])
code(s, L, I(4.2), I(6.6), I(2.0), """
int x = 10;
int *p = &x;      // p trỏ tới x
*p = 20;          // sửa x thông qua p
printf("%d", x);  // 20
""", 17)
text(s, I(7.8), TOP, I(5), I(0.4), "p trỏ tới x", 18, True, BLUE, PP_ALIGN.CENTER)
for x, name, val, addr, fill, sz, f in [
        (8.0, "p", "0x1004", "địa chỉ: 0x2000", WHITE, 18, MONO),
        (10.8, "x", "20", "địa chỉ: 0x1004", BLUE_PALE, 24, FONT)]:
    b = rect(s, I(x), I(2.9), I(1.9), I(1.0), fill, line=BLUE)
    label(b, val, sz, True, font=f)
    text(s, I(x), I(2.4), I(1.9), I(0.4), name, 20, True, align=PP_ALIGN.CENTER)
    text(s, I(x), I(3.95), I(1.9), I(0.4), addr, 13, color=GRAY, align=PP_ALIGN.CENTER)
arrow(s, I(9.9), I(3.4), I(10.8), I(3.4))
table(s, I(8.0), I(4.75), I(4.7), [
    ["Biểu thức", "Giá trị"],
    ["`p`", "0x1004 (địa chỉ của x)"],
    ["`*p`", "20 (giá trị của x)"],
    ["`&p`", "0x2000 (địa chỉ của p)"],
], [1.4, 3.3], 15, 0.4)
note(s, "Dấu * khi khai báo khác dấu * trong biểu thức. sizeof(con trỏ) = 8 byte trên máy 64-bit.")

# ===== Slide 5 – Con trỏ và mảng =====
s = slide("3. Con trỏ và mảng", 5)
bullets(s, L, TOP, I(6.6), I(2.8), [
    "Tên mảng ≈ địa chỉ phần tử đầu: `a == &a[0]`",
    "`a[i]` tương đương `*(a + i)`",
    "Số học con trỏ: `p + 1` tiến thêm `sizeof(kiểu)` byte",
    "Phần tử mảng nằm **liên tiếp** trong bộ nhớ",
])
code(s, L, I(4.3), I(6.6), I(1.9), """
int a[5] = {1, 2, 3, 4, 5};
int *p = a;
for (int i = 0; i < 5; i++)
    printf("%d ", *(p + i));
""", 17)
text(s, I(7.6), TOP, I(5.3), I(0.4), "Mảng int a[5] trong bộ nhớ", 18, True, BLUE,
     PP_ALIGN.CENTER)
for i in range(5):
    x = I(7.75 + i * 1.0)
    b = rect(s, x, I(2.6), I(1.0), I(0.9), BLUE_PALE if i == 0 else WHITE, line=BLUE)
    label(b, str(i + 1), 22, True)
    text(s, x, I(3.55), I(1.0), I(0.4), str(1000 + 4 * i), 13, color=GRAY, font=MONO,
         align=PP_ALIGN.CENTER)
    text(s, x, I(2.05), I(1.0), I(0.4), "p" if i == 0 else f"p+{i}", 16, True,
         BLUE_LIGHT, PP_ALIGN.CENTER, font=MONO)
text(s, I(7.6), I(4.0), I(5.3), I(0.4), "mỗi bước +4 byte = sizeof(int)", 14,
     color=GRAY, align=PP_ALIGN.CENTER)
table(s, I(7.75), I(4.65), I(5.0), [
    ["", "Mảng a", "Con trỏ p"],
    ["sizeof", "20 byte", "8 byte"],
    ["Gán lại", "Không", "Được"],
], [1.4, 1.8, 1.8], 15, 0.42)
note(s, "Mảng không phải con trỏ: sizeof khác nhau, không gán lại tên mảng. C không kiểm tra biên.")

# ===== Slide 6 – Con trỏ và hàm =====
s = slide("4. Con trỏ và hàm", 6)
bullets(s, L, TOP, FULL_W, I(1.3), [
    "C truyền tham số **theo giá trị**: hàm chỉ nhận **bản sao**",
    "Muốn hàm sửa được biến gốc → truyền **địa chỉ**",
])
text(s, L, I(2.75), I(5.9), I(0.4), "✗  Truyền giá trị – x, y không đổi", 18, True, RED)
code(s, L, I(3.2), I(5.9), I(1.9), """
void swap(int a, int b) {
    int t = a;
    a = b;
    b = t;
}
""", 16)
text(s, I(6.85), I(2.75), I(5.9), I(0.4), "✓  Truyền địa chỉ – hoán đổi đúng", 18, True, GREEN)
code(s, I(6.85), I(3.2), I(5.9), I(1.9), """
void swap(int *a, int *b) {
    int t = *a;
    *a = *b;
    *b = t;
}   // gọi: swap(&x, &y);
""", 16)
bullets(s, L, I(5.35), FULL_W, I(1.6), [
    "Ứng dụng: trả về nhiều kết quả, truyền mảng lớn không sao chép, `scanf(\"%d\", &x)`",
    "**Con trỏ hàm**: truyền hàm làm tham số (callback) – `qsort`, `pthread_create`",
], 19, 6)
note(s, "Bản chất con trỏ vẫn truyền theo giá trị (bản sao địa chỉ). Không trả về địa chỉ biến cục bộ.")

# ===== Slide 7 – Cấp phát động =====
s = slide("5. Cấp phát bộ nhớ động", 7)
bullets(s, L, TOP, I(6.4), I(2.4), [
    "Dùng khi **lúc chạy** mới biết kích thước",
    "Bộ nhớ nằm trên **Heap**, trả về một **con trỏ**",
    "**Phải tự giải phóng**, nếu không sẽ rò rỉ bộ nhớ",
    "Luôn kiểm tra kết quả cấp phát (`NULL`)",
], 21)
table(s, L, I(4.25), I(6.4), [
    ["", "C", "C++"],
    ["Cấp phát", "`malloc`, `calloc`", "`new`, `new[]`"],
    ["Giải phóng", "`free`", "`delete`, `delete[]`"],
    ["Khi lỗi", "trả về NULL", "ném bad_alloc"],
], [1.6, 2.4, 2.4], 16, 0.48)
text(s, I(7.3), TOP, I(5.4), I(0.4), "C", 18, True, BLUE)
code(s, I(7.3), I(1.9), I(5.4), I(2.3), """
int n = 100;
int *arr = (int*)malloc(n * sizeof(int));
if (arr == NULL) return 1;
// ... sử dụng arr ...
free(arr);
arr = NULL;
""", 15)
text(s, I(7.3), I(4.4), I(5.4), I(0.4), "C++", 18, True, BLUE)
code(s, I(7.3), I(4.85), I(5.4), I(1.3), """
int *arr = new int[n];
// ... sử dụng arr ...
delete[] arr;
""", 15)
note(s, "Stack nhỏ, tự giải phóng; Heap lớn, tự quản lý. new[] phải đi với delete[].")

# ===== Slide 8 – Lỗi thường gặp =====
s = slide("6. Lỗi thường gặp với con trỏ", 8)
errs = [("Wild pointer", "Chưa khởi tạo,\ntrỏ tới địa chỉ ngẫu nhiên", "int *p;\n*p = 5;"),
        ("NULL pointer", "Truy cập *p khi p == NULL\n→ Segmentation fault", "int *p = NULL;\n*p = 5;"),
        ("Dangling pointer", "Dùng con trỏ sau khi\nđã free / delete", "free(p);\n*p = 5;"),
        ("Memory leak", "Cấp phát nhưng\nkhông giải phóng", "p = malloc(..);\np = malloc(..);")]
cw = I(2.9)
for i, (name, desc, src) in enumerate(errs):
    x = L + i * (cw + I(0.18))
    label(rect(s, x, TOP, cw, I(0.6), BLUE), name, 19, True, WHITE)
    label(rect(s, x, TOP + I(0.6), cw, I(1.2), WHITE, line=BLUE), desc, 15)
    code(s, x, TOP + I(1.9), cw, I(0.95), src, 14)
tip = rect(s, L, I(4.75), FULL_W, I(2.0), BLUE_PALE, line=BLUE_LIGHT)
tip.text_frame.margin_left = I(0.3)
tip.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
for i, ln in enumerate(["**Cách phòng tránh**",
                        "•  Luôn khởi tạo: `int *p = NULL;`  (C++: `nullptr`)",
                        "•  Sau `free(p);` gán `p = NULL;` – mỗi `malloc` đi kèm một `free`",
                        "•  Công cụ kiểm tra: Valgrind, AddressSanitizer (`-fsanitize=address`)"]):
    p = tip.text_frame.paragraphs[0] if i == 0 else tip.text_frame.add_paragraph()
    p.space_after = Pt(4)
    rich(p, ln, 19, BLUE if i == 0 else BLACK)
note(s, "Dangling pointer nguy hiểm nhất vì chương trình có thể vẫn chạy 'bình thường'.")

# ===== Slide 9 – C++ hiện đại =====
s = slide("7. Con trỏ trong C++ hiện đại", 9)
bullets(s, L, TOP, I(6.3), I(4.8), [
    "**Tham chiếu** (`&`): tên gọi khác của biến, không NULL, không đổi đối tượng",
    "**Smart pointer** (`<memory>`): tự giải phóng khi ra khỏi phạm vi (RAII)",
    ("-", "`unique_ptr`: một chủ sở hữu, không sao chép được"),
    ("-", "`shared_ptr`: nhiều chủ sở hữu, đếm tham chiếu, về 0 thì tự xóa"),
    ("-", "`weak_ptr`: chỉ quan sát, tránh vòng tham chiếu"),
    "Ưu tiên `make_unique` / `make_shared` thay vì `new`",
], 21, 8)
text(s, I(7.2), TOP, I(5.5), I(0.4), "Tham chiếu", 18, True, BLUE)
code(s, I(7.2), I(1.9), I(5.5), I(1.3), """
void swap(int &a, int &b) {
    int t = a; a = b; b = t;
}
swap(x, y);   // không cần &x, &y
""", 15)
text(s, I(7.2), I(3.4), I(5.5), I(0.4), "Smart pointer", 18, True, BLUE)
code(s, I(7.2), I(3.85), I(5.5), I(2.3), """
#include <memory>
auto p = std::make_unique<int>(10);
auto q = std::make_shared<int>(20);
auto r = q;   // đếm tham chiếu = 2
// không cần delete
""", 15)
note(s, "Bộ đếm của shared_ptr là atomic – an toàn đa luồng; dữ liệu trỏ tới thì không.")

# ===== Slide 10 – Song song =====
s = slide("8. Con trỏ trong tính toán song song", 10)
bullets(s, L, TOP, I(6.4), I(3.0), [
    "Các luồng (thread) **dùng chung bộ nhớ** của tiến trình",
    "Con trỏ giúp **chia dữ liệu** cho từng luồng, **không sao chép**",
    "`pthread_create` truyền tham số qua `void *`",
    "Nhiều tham số → gói vào `struct`, truyền địa chỉ struct",
], 21)
code(s, L, I(4.55), I(6.4), I(1.6), """
int pthread_create(pthread_t *thread,
        const pthread_attr_t *attr,
        void *(*start_routine)(void *),  // con trỏ hàm
        void *arg);                      // con trỏ void
""", 13)
warn = rect(s, I(7.4), TOP, I(5.3), I(4.6), ORANGE_PALE, line=ORANGE)
tf = warn.text_frame
tf.word_wrap = True
tf.vertical_anchor = MSO_ANCHOR.TOP
tf.margin_left, tf.margin_right, tf.margin_top = I(0.25), I(0.2), I(0.15)
for i, (t, sz) in enumerate([
        ("⚠  Cần chú ý", 20),
        ("**Race condition**: nhiều luồng cùng ghi một địa chỉ → kết quả sai", 17),
        ("    → dùng mutex / atomic, hoặc mỗi luồng ghi vùng riêng", 16),
        ("**Địa chỉ không hợp lệ**: truyền `&i` của vòng lặp, hoặc biến cục bộ đã hết phạm vi", 17),
        ("    → mỗi luồng một phần tử riêng: `&tk[i]`", 16),
        ("**False sharing**: biến khác nhau nằm chung cache line → chậm", 17)]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(8)
    rich(p, t, sz, ORANGE if i == 0 else BLACK, bold=(i == 0))
text(s, I(7.4), I(6.15), I(5.3), I(0.6),
     "Liên hệ: OpenMP, CUDA (cudaMalloc), MPI (con trỏ tới bộ đệm)", 14, color=GRAY)
note(s, "Luồng chung Heap/biến toàn cục, mỗi luồng có Stack riêng. void* phải ép kiểu trước khi dùng.")

# ===== Slide 11 – Ví dụ =====
s = slide("9. Ví dụ: Tính tổng mảng bằng nhiều luồng", 11)
code(s, L, TOP, I(7.0), I(5.35), """
typedef struct { int *start; int len; long sum; } Task;

void *work(void *arg) {
    Task *t = (Task *)arg;           // ép void* về Task*
    t->sum = 0;
    for (int i = 0; i < t->len; i++)
        t->sum += *(t->start + i);
    return NULL;
}

// trong main():  N = 8, T = 2
for (int i = 0; i < T; i++) {
    tk[i].start = a + i * (N / T);   // số học con trỏ
    tk[i].len   = N / T;
    pthread_create(&th[i], NULL, work, &tk[i]);
}
for (int i = 0; i < T; i++) {
    pthread_join(th[i], NULL);
    total += tk[i].sum;              // total = 36
}
""", 14)
text(s, I(7.9), TOP, I(4.9), I(0.4), "Chia mảng cho 2 luồng", 18, True, BLUE, PP_ALIGN.CENTER)
for i in range(8):
    b = rect(s, I(8.0 + i * 0.58), I(2.55), I(0.58), I(0.7),
             BLUE_PALE if i < 4 else ORANGE_PALE, line=BLUE)
    label(b, str(i + 1), 18, True)
text(s, I(8.0), I(2.05), I(2.32), I(0.4), "a + 0", 15, True, BLUE, PP_ALIGN.CENTER, font=MONO)
text(s, I(10.32), I(2.05), I(2.32), I(0.4), "a + 4", 15, True, ORANGE, PP_ALIGN.CENTER, font=MONO)
label(rect(s, I(8.0), I(3.5), I(2.25), I(0.75), BLUE_PALE, line=BLUE), "Luồng 0\nsum = 10", 15, True)
label(rect(s, I(10.4), I(3.5), I(2.25), I(0.75), ORANGE_PALE, line=BLUE), "Luồng 1\nsum = 26", 15, True)
label(rect(s, I(9.2), I(4.6), I(2.25), I(0.6), BLUE), "Tổng = 36", 18, True, WHITE)
arrow(s, I(9.12), I(4.25), I(9.8), I(4.6))
arrow(s, I(11.52), I(4.25), I(10.85), I(4.6))
bullets(s, I(7.9), I(5.45), I(5.0), I(1.5), [
    "Không sao chép mảng",
    "Mỗi luồng ghi `sum` riêng → không race",
], 17, 4)
note(s, "Biên dịch: gcc sum_pthread.c -o sum -pthread. Code đầy đủ: code/sum_pthread.c")

# ===== Slide 12 – Kết luận =====
s = slide("Kết luận", 12)
for i, (k, v) in enumerate([
        ("Mạnh mẽ", "Thao tác trực tiếp với bộ nhớ, hiệu năng cao"),
        ("Nền tảng", "Mảng, hàm, cấp phát động, cấu trúc dữ liệu"),
        ("Song song", "Chia sẻ dữ liệu giữa các luồng, không cần sao chép"),
        ("Cẩn thận", "Dễ lỗi bộ nhớ & race condition – C++ ưu tiên smart pointer")]):
    y = I(1.6 + i * 1.3)
    label(rect(s, I(1.2), y, I(2.6), I(1.0), BLUE), k, 22, True, WHITE)
    vb = rect(s, I(3.8), y, I(8.3), I(1.0), BLUE_PALE)
    label(vb, v, 22, align=PP_ALIGN.LEFT)
    vb.text_frame.margin_left = I(0.3)
note(s, "Tóm tắt 4 ý chính.")

# ===== Slide 13 – Cảm ơn =====
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, SW, I(0.25), BLUE)
text(s, 0, I(2.2), SW, I(1.4), "CẢM ƠN THẦY/CÔ VÀ CÁC BẠN\nĐÃ LẮNG NGHE!", 40, True,
     BLUE, PP_ALIGN.CENTER)
rect(s, I(5.17), I(4.0), I(3.0), Pt(3), BLUE_LIGHT)
text(s, 0, I(4.3), SW, I(0.6), "Hỏi & Đáp", 28, align=PP_ALIGN.CENTER)
text(s, 0, I(5.4), SW, I(0.5), "Bùi Xuân Chung  •  Lê Văn Cường  •  Đào Ngọc Cường", 18,
     color=GRAY, align=PP_ALIGN.CENTER)
rect(s, 0, SH - I(0.25), SW, I(0.25), BLUE)

# ===== Lưu =====
root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
out_dir = os.path.join(root, "slides")
os.makedirs(out_dir, exist_ok=True)
out = os.path.join(out_dir, "ConTro_C_CPP.pptx")
prs.save(out)
print("Saved:", out, "-", len(prs.slides), "slides")
